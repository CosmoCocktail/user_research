import streamlit as st
import pandas as pd
import ast
import os
import random
from openpyxl import load_workbook, Workbook

# ======================
# 페이지 설정
# ======================
st.set_page_config(
    page_title="팀플 빌런즈 : 눈 속 마을 빌런 테스트",
    page_icon="🐧",
    layout="centered"
)

# ======================
# 상수
# ======================
QUESTION_FILE = "questions.xlsx"
RESULT_FILE   = "result.xlsx"
CHARACTERS    = ["에디", "크롱", "뽀로로", "루피", "포비"]



# ======================
# 동점 추가 질문
# 동점 캐릭터에 해당하는 답변만 선택지로 표시
# ======================
TIEBREAK_QUESTION = {
    "질문": "주제를 정하는데 의견이 갈린다면?",
    "에디":   "\"잠깐, 지금 의견 말고 나 완전 좋은 생각 났어.\"",
    "크롱":   "\"다 괜찮은 것 같은데…\"",
    "뽀로로": "\"제일 재밌어 보이는 거 하면 안 돼?\"",
    "루피":   "\"일단 기준부터 정하고 제일 괜찮은 안으로 가자.\"",
    "포비":   "\"다들 말해봐. 내가 의견을 정리해볼게.\"",
}

# ======================
# 데이터 로드
# ======================
@st.cache_data
def load_questions():
    df = pd.read_excel(QUESTION_FILE)
    df = df.dropna(subset=["질문", "답변A", "답변B"])
    df = df[df["질문"].astype(str).str.strip() != ""]
    df = df[df["답변A"].astype(str).str.strip().str.lower() != "nan"]
    df = df[df["답변B"].astype(str).str.strip().str.lower() != "nan"]
    df = df.reset_index(drop=True)
    return df

questions_df = load_questions()

# ======================
# 점수 파싱 함수
# ======================
def parse_types(type_str: str) -> list:
    try:
        result = ast.literal_eval(str(type_str).strip())
        return [r.strip() for r in result]
    except Exception:
        return []

# ======================
# result.xlsx — 해당 캐릭터 열 +1 갱신
# 구조: 행1=헤더(캐릭터명), 행2=누적 카운트
# ======================
def save_result(character: str):
    if os.path.exists(RESULT_FILE):
        wb = load_workbook(RESULT_FILE)
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        if character in header:
            col_idx = header.index(character) + 1
            cur = ws.cell(row=2, column=col_idx).value or 0
            ws.cell(row=2, column=col_idx).value = cur + 1
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(CHARACTERS)
        ws.append([1 if c == character else 0 for c in CHARACTERS])
    wb.save(RESULT_FILE)

# ======================
# result.xlsx — 집계 읽기
# ======================
def load_result_counts():
    if not os.path.exists(RESULT_FILE):
        return {c: 0 for c in CHARACTERS}, 0
    df = pd.read_excel(RESULT_FILE, header=0)
    if not all(c in df.columns for c in CHARACTERS):
        return {c: 0 for c in CHARACTERS}, 0
    counts = {c: int(df[c].iloc[0]) for c in CHARACTERS}
    return counts, sum(counts.values())

# ======================
# 점수 평가 → 결과 or 동점 추가 질문 분기
# ======================
def evaluate_and_route():
    scores    = st.session_state.scores
    max_score = max(scores.values())
    top       = [c for c, s in scores.items() if s == max_score]

    if len(top) == 1:
        finalize(top[0])
    else:
        st.session_state.tied_chars = top
        st.session_state.tb_order   = None
        st.session_state.page       = "tiebreak"
    st.rerun()

def finalize(character: str):
    st.session_state.result_character = character
    if not st.session_state.result_saved:
        save_result(character)
        st.session_state.result_saved = True
    st.session_state.page = "result"

# ======================
# 세션 초기화
# ======================
defaults = {
    "page":             "home",
    "question_idx":     0,
    "scores":           {c: 0 for c in CHARACTERS},
    "result_character": None,
    "result_saved":     False,
    "tied_chars":       [],
    "tb_order":         None,
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v

def reset_session():
    for k, v in defaults.items():
        st.session_state[k] = v if k != "scores" else {c: 0 for c in CHARACTERS}
    for k in list(st.session_state.keys()):
        if k.startswith("q_order_"):
            del st.session_state[k]

# ======================
# 공통 스타일
# ======================
st.markdown("""
<style>
    .big-title  { text-align:center; font-size:2rem; font-weight:800; margin-bottom:.5rem; }
    .sub-title  { text-align:center; font-size:1.1rem; color:#666; margin-bottom:1.5rem; }
    .q-text     { text-align:center; font-size:1.35rem; font-weight:700; margin-bottom:1.2rem; line-height:1.5; }
    .prog-text  { text-align:center; font-size:.9rem; color:#888; margin-bottom:.4rem; }
    .tie-badge  { display:inline-block; background:#ff6b35; color:#fff;
                  border-radius:20px; padding:4px 18px; font-size:.9rem; margin-bottom:1rem; }
    .res-char   { text-align:center; font-size:2.4rem; font-weight:900; margin:1rem 0; }
    .stat-box   { background:#f0f4ff; border-radius:12px; padding:1.2rem; margin-top:1rem; }
    .stat-total { text-align:center; font-size:1.6rem; font-weight:800; color:#3355ff; }
    .char-tag   { display:inline-block; border-radius:14px; padding:3px 12px;
                  font-size:.78rem; font-weight:700; margin-bottom:4px; }
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════
# 홈 페이지
# ══════════════════════════════════════════
if st.session_state.page == "home":

    st.markdown("<div class='big-title'>팀플 빌런즈</div>", unsafe_allow_html=True)
    st.markdown("<div class='sub-title'> 눈 속 마을 빌런 테스트</div>", unsafe_allow_html=True)

    if os.path.exists("단체컷.png"):
        st.image("단체컷.png", use_container_width=True)
    st.write("")

    if st.button("🐧 시작하기", use_container_width=True):
        reset_session()
        st.session_state.page = "question"
        st.rerun()


# ══════════════════════════════════════════
# 질문 페이지 (A/B 2지선다, questions.xlsx 기반)
# ══════════════════════════════════════════
elif st.session_state.page == "question":

    total = len(questions_df)
    idx   = st.session_state.question_idx

    # 모든 질문 완료 → 점수 평가 후 즉시 중단
    if idx >= total:
        evaluate_and_route()
        st.stop()

    row = questions_df.iloc[idx]

    # 진행 상황
    st.markdown(f"<div class='prog-text'>질문 {idx+1} / {total}</div>", unsafe_allow_html=True)
    st.progress(idx / total)
    st.write("")

    # 질문 텍스트 (엑셀 질문 항목) — 이미지 위에 항상 출력
    st.markdown(f"<div class='q-text'>{row['질문']}</div>", unsafe_allow_html=True)

    # 질문 이미지 중앙 출력 (question1.png ~ question11.png)
    img_path = f"question{idx+1}.png"
    if os.path.exists(img_path):
        try:
            col_l, col_c, col_r = st.columns([1, 3, 1])
            with col_c:
                st.image(img_path, use_container_width=True)
        except Exception:
            pass  # 이미지 손상/형식 오류 시 텍스트만 표시
    st.write("")

    # A/B 답변 — 리렌더링 시 순서 유지
    order_key = f"q_order_{idx}"
    if order_key not in st.session_state:
        opts = [
            ("A", str(row["답변A"]), str(row["답변A_유형"]),
             int(float(row["답변A_점수"])) if str(row["답변A_점수"]) not in ("", "nan") else 2),
            ("B", str(row["답변B"]), str(row["답변B_유형"]),
             int(float(row["답변B_점수"])) if str(row["답변B_점수"]) not in ("", "nan") else 2),
        ]
        random.shuffle(opts)
        st.session_state[order_key] = opts
    opts = st.session_state[order_key]

    # 각 답변에 해당하는 캐릭터 변수 태그 추출
    def get_tags_html(types_str: str) -> str:
        chars = parse_types(types_str)
        tags  = []
        for c in chars:
            info  = CHAR_INFO.get(c, {})
            color = info.get("color", "#888")
            emoji = info.get("emoji", "")
            tag   = info.get("tag", c)
            tags.append(
                f"<span class='char-tag' style='background:{color};color:#fff'>"
                f"{emoji} {c} · {tag}</span>"
            )
        return " ".join(tags)

    col1, col2 = st.columns(2)
    for col, (label, text, types_str, score) in zip([col1, col2], opts):
        with col:
            if not text or text.strip() in ("", "nan"):
                continue
            # 캐릭터 변수 태그 표시
            tags_html = get_tags_html(types_str)
            st.markdown(f"<div style='text-align:center;margin-bottom:6px'>{tags_html}</div>",
                        unsafe_allow_html=True)
            if st.button(text, use_container_width=True, key=f"q{idx}_{label}"):
                for char in parse_types(types_str):
                    if char in st.session_state.scores:
                        st.session_state.scores[char] += score
                st.session_state.question_idx += 1
                st.rerun()


# ══════════════════════════════════════════
# 동점 추가 질문 페이지
# 동점 캐릭터에 해당하는 답변만 선택지로 표시
# ══════════════════════════════════════════
elif st.session_state.page == "tiebreak":

    tied     = st.session_state.tied_chars
    tied_str = ", ".join(tied)

    # 동점 캐릭터 태그 표시
    tags_html = " ".join([
        "<span class='char-tag' style='background:" + CHAR_INFO[c]["color"] + ";color:#fff'>"
        + CHAR_INFO[c]["emoji"] + " " + c + "</span>"
        for c in tied if c in CHAR_INFO
    ])
    st.markdown(
        f"<div style='text-align:center;margin-bottom:.5rem'>{tags_html}</div>",
        unsafe_allow_html=True
    )
    st.markdown(
        f"<div style='text-align:center'>"
        f"<span class='tie-badge'>⚖️ 두구두구! 마지막 질문</span>"
        f"</div>",
        unsafe_allow_html=True
    )
    st.write("")

    # 동점 추가 질문 텍스트
    st.markdown(
        f"<div class='q-text'>{TIEBREAK_QUESTION['질문']}</div>",
        unsafe_allow_html=True
    )
    st.write("")

    # 동점 캐릭터 답변만 선택지 표시, 순서 고정
    if st.session_state.tb_order is None:
        opts = [(char, TIEBREAK_QUESTION[char]) for char in tied if char in TIEBREAK_QUESTION]
        random.shuffle(opts)
        st.session_state.tb_order = opts
    opts = st.session_state.tb_order

    col_list = st.columns(len(opts))
    for col, (char, text) in zip(col_list, opts):
        with col:
            info = CHAR_INFO.get(char, {})
            # 각 선택지 위에 해당 캐릭터 변수 태그 표시
            st.markdown(
                f"<div style='text-align:center;margin-bottom:6px'>"
                f"<span class='char-tag' style='background:{info.get('color','#888')};color:#fff'>"
                f"{info.get('emoji','')} {char} · {info.get('tag','')}</span></div>",
                unsafe_allow_html=True
            )
            if st.button(text, use_container_width=True, key=f"tb_{char}"):
                finalize(char)
                st.rerun()


# ══════════════════════════════════════════
# 결과 페이지
# ══════════════════════════════════════════
elif st.session_state.page == "result":

    # result_character가 None이면 홈으로 복귀 (비정상 진입 방어)
    if not st.session_state.result_character:
        st.session_state.page = "home"
        st.rerun()

    character     = st.session_state.result_character
    info          = CHAR_INFO.get(character, {})
    counts, total = load_result_counts()

    st.markdown("<div class='big-title'>🎉 결과 발표!</div>", unsafe_allow_html=True)
    st.write("")

    # 결과 캐릭터 이미지
    img_path = f"{character}.png"
    if os.path.exists(img_path):
        st.image(img_path, use_container_width=True)

    # 캐릭터 변수 태그
    st.markdown(
        f"<div style='text-align:center;margin-bottom:.5rem'>"
        f"<span class='char-tag' style='background:{info.get('color','#888')};color:#fff;font-size:1rem;padding:5px 18px'>"
        f"{info.get('emoji','')} {character} · {info.get('tag','')}</span></div>",
        unsafe_allow_html=True
    )
    st.markdown(
        f"<div class='res-char'>당신은 <span style='color:{info.get('color','#3355ff')}'>{character}</span>입니다!</div>",
        unsafe_allow_html=True
    )
    st.write("")

    # 참여 통계
    st.markdown("<div class='stat-box'>", unsafe_allow_html=True)
    st.markdown(
        "<div style='text-align:center;font-size:1rem;color:#444;margin-bottom:.5rem'>"
        "지금까지 이 테스트에 참여한 인원</div>",
        unsafe_allow_html=True
    )
    st.markdown(f"<div class='stat-total'>총 {total}명</div>", unsafe_allow_html=True)
    st.write("")

    st.markdown("**캐릭터별 결과 분포**")
    for char in CHARACTERS:
        cinfo = CHAR_INFO.get(char, {})
        count = counts.get(char, 0)
        pct   = round(count / total * 100, 1) if total > 0 else 0.0
        st.markdown(
            f"<span class='char-tag' style='background:{cinfo.get('color','#888')};color:#fff'>"
            f"{cinfo.get('emoji','')} {char}</span> "
            f"**{pct}%** 의 인원이 **{char}** 를 선택했습니다. ({count}명)",
            unsafe_allow_html=True
        )

    st.markdown("</div>", unsafe_allow_html=True)
    st.write("")

    if st.button("🔄 다시하기", use_container_width=True):
        reset_session()
        st.rerun()
